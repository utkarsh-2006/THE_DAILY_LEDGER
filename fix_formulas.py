import openpyxl

wb = openpyxl.load_workbook('THE_DAILY_LEDGER_SRS.xlsx')
ws = wb['13_FORMULAS']

formula_map = {
    "Pre-District Current Yield": "Base Yield\n× (1.0 + (0.50 × Num_Primary_Cashflow) + (0.35 × Num_Secondary_Cashflow))\n× Market Yield Modifier",
    "Final Current Yield": "Pre-District Current Yield\n× (1.0 + District_Control_Modifier + Network_Modifier)"
}

for r in range(2, ws.max_row + 1):
    name = str(ws.cell(row=r, column=1).value)
    if name in formula_map:
        ws.cell(row=r, column=2).value = formula_map[name]

wb.save('THE_DAILY_LEDGER_SRS.xlsx')
print("Formulas updated.")
