import codecs
import openpyxl

# --- Part 1: Update Markdown ---
try:
    with codecs.open('SRS_THE_DAILY_LEDGER.md', 'r', 'utf-8') as f:
        md = f.read()

    md = md.replace(
        "Turn Loop (Per Player): Move → Identify Space → Check Forced State → Resolve Special Space → Resolve Mandatory Payment → Optional Actions → End Turn.",
        "Turn Loop (Per Player): Move → Identify Space → Check Forced State → Resolve Property or Special Space → Resolve Mandatory Payment → Optional Actions → End Turn."
    )
    md = md.replace("RESOLVE_SPECIAL_SPACE", "RESOLVE_PROPERTY_OR_SPECIAL_SPACE")

    with codecs.open('SRS_THE_DAILY_LEDGER.md', 'w', 'utf-8') as f:
        f.write(md)
    print("Markdown updated.")
except Exception as e:
    print("MD Update Error:", e)

# --- Part 2: Update Excel ---
try:
    wb = openpyxl.load_workbook('THE_DAILY_LEDGER_SRS.xlsx')

    # 1. Update Test Cases Status
    if '17_TEST_CASES' in wb.sheetnames:
        ws_tc = wb['17_TEST_CASES']
        for r in range(2, ws_tc.max_row + 1):
            tc_id = str(ws_tc.cell(row=r, column=1).value)
            if tc_id.startswith("TC-02") or tc_id.startswith("TC-03"):
                try:
                    num = int(tc_id.split("-")[1])
                    if 25 <= num <= 36:
                        ws_tc.cell(row=r, column=3).value = "NOT_RUN"
                except:
                    pass

    # 2. Strengthen Acceptance Criteria
    ac_map = {
        "AC-003": "Property resolution occurs before OPTIONAL_ACTIONS during every normal player turn.",
        "AC-004": "When a player lands on a property owned by another player, the applicable rent obligation is resolved before any optional development or trade action.",
        "AC-005": "When a player lands on an unowned property, BUY or PASS/Auction is resolved during RESOLVE_PROPERTY_OR_SPECIAL_SPACE; the property decision does not reappear as a later optional action.",
        "AC-006": "Final Current Yield is calculated by applying District Control and Development Network modifiers after the Pre-District Current Yield calculation.",
        "AC-007": "District Control contributes exactly +5% Current Yield when the player controls the district and 0% otherwise.",
        "AC-008": "An active Development Network contributes exactly +5% Current Yield when its eligibility requirements are satisfied and 0% otherwise.",
        "AC-009": "When bankruptcy reduces the number of active players to exactly one, the game enters Endgame immediately after required elimination and state cleanup.",
        "AC-010": "After an early Endgame trigger caused by bankruptcy, no additional player turn or Market Pulse is resolved.",
        "AC-011": "Civic Reserve remains ACTIVE across round boundaries until claimed and becomes INACTIVE immediately after the first valid claim.",
        "AC-012": "Treasury Window eligibility is true only when the player's current cash is strictly below $300 at the moment the Treasury Window landing resolves.",
        "AC-013": "Transport relocation bypasses intermediate spaces and does not award the Velora Central dividend or trigger destination landing effects.",
        "AC-014": "Forced entry into Civic Hold applies LOCKED status and causes the player to miss exactly one complete normal turn."
    }

    if '18_ACCEPTANCE' in wb.sheetnames:
        ws_ac = wb['18_ACCEPTANCE']
        for r in range(2, ws_ac.max_row + 1):
            ac_id = str(ws_ac.cell(row=r, column=1).value)
            if ac_id in ac_map:
                ws_ac.cell(row=r, column=2).value = ac_map[ac_id]

    # 3. Global Term Audit & Replace
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                if isinstance(cell.value, str):
                    val = cell.value
                    if "RESOLVE_SPECIAL_SPACE" in val:
                        val = val.replace("RESOLVE_SPECIAL_SPACE", "RESOLVE_PROPERTY_OR_SPECIAL_SPACE")
                    cell.value = val

    wb.save('THE_DAILY_LEDGER_SRS.xlsx')
    print("Excel updated.")
except Exception as e:
    print("Excel Update Error:", e)
