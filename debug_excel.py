import openpyxl
wb = openpyxl.load_workbook('THE_DAILY_LEDGER_SRS.xlsx')
ws = wb['03_BOARD_SPACES']
print("Max Col:", ws.max_column)
headers = {cell.value: idx for idx, cell in enumerate(ws[1])}
print("Headers:", headers)
